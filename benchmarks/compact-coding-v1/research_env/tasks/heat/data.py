"""Figshare CORE v1 importer; sheet/cell lineage and prefix-only wearable inputs."""
from pathlib import Path
import hashlib
from research_env.core.common import file_hash, finite
from research_env.tasks.tabular import export_dataset

SOURCE = 'https://doi.org/10.6084/m9.figshare.29120024.v1'
MD5 = '01214bfc835215fe1a5d3f5da1e9ad52'


def minute(value):
    if isinstance(value, str) and '-' in value:
        return int(value.split('-')[1])  # interval average available at its right edge
    return int(value)


def make_request(history, demographics, cutoff):
    return {'history': [{k: row[k] for k in ('time_min', 'hr_bpm', 'skin_c', 'ambient_c', 'humidity_pct')}
                        for row in history if row['time_min'] <= cutoff],
            'demographics': {k: demographics[k] for k in ('age_years', 'mass_kg', 'height_cm')},
            'cutoff_min': cutoff, 'horizons_min': [5, 10, 20], 'state': None}


def prepare(source_root, output_root, split_manifest=None):
    import openpyxl
    path = Path(source_root)/'core_cycling_2025.xlsx'
    if hashlib.md5(path.read_bytes()).hexdigest() != MD5:
        raise ValueError('Original Figshare workbook checksum mismatch')
    w = openpyxl.load_workbook(path, data_only=True)
    people, records, flags, lineage = [], [], [], []
    core, skin, exercise, ambient = (w[n] for n in ('Core Temp', 'Skin Temp', 'Exercise Data', 'Environmental Conditions'))
    core_cols = {int(core.cell(2, c).value): c for c in range(3, 27)}
    skin_cols = {int(skin.cell(2, c).value): c for c in range(3, 27)}
    hr_cols = {int(exercise.cell(2, c).value): c for c in range(28, 52)}
    env_cols = {int(ambient.cell(2, c).value): c for c in range(3, 27)}
    for drow in range(2, w['Demographics'].max_row+1):
        d = [w['Demographics'].cell(drow, c).value for c in range(1, 6)]
        if not isinstance(d[0], int):
            continue
        subject = d[0]
        person = f'H{subject:02}'
        demo = dict(zip(('age_years', 'height_cm', 'mass_kg'), map(float, d[2:])))
        if not (18 <= demo['age_years'] <= 100 and 100 <= demo['height_cm'] <= 230 and 30 <= demo['mass_kg'] <= 250):
            raise ValueError('Invalid anthropometry')
        y, skin_by_t, hr_by_t, env_by_t = {}, {}, {}, {}
        for sheet, col, target, lower, upper in (
                (core, core_cols[subject], y, 30, 45),
                (skin, skin_cols[subject], skin_by_t, 20, 45),
                (exercise, hr_cols[subject], hr_by_t, 30, 240)):
            for row in range(3, sheet.max_row+1):
                t, v = sheet.cell(row, 2).value, sheet.cell(row, col).value
                if t is None:
                    continue
                t = minute(t)
                if v is None:
                    continue
                cell = f'{sheet.title}!{sheet.cell(row,col).coordinate}'
                if not finite(v) or not lower <= v <= upper:
                    flags.append({'person': person, 'cell': cell, 'original': v,
                                  'rule': f'omit outside [{lower},{upper}]; no imputation'})
                    continue
                if t in target:
                    raise ValueError('Duplicate sheet timestamp')
                target[t] = float(v)
                lineage.append({'person': person, 'time_min': t, 'cell': cell, 'value': v})
        for row in range(3, ambient.max_row+1):
            t = minute(ambient.cell(row, 2).value)
            col = env_cols[subject]
            values = [ambient.cell(row, col).value, ambient.cell(row, col+25).value]
            if all(finite(v) for v in values) and 0 <= values[0] <= 60 and 0 <= values[1] <= 100:
                env_by_t[t] = list(map(float, values))
                for c, v in zip((col, col+25), values):
                    lineage.append({'person': person, 'time_min': t, 'cell': f'{ambient.title}!{ambient.cell(row,c).coordinate}', 'value': v})
        history = []
        for t in sorted(skin_by_t):
            # No future interpolation. HR/environment carry forward only after observation.
            hrs = [k for k in hr_by_t if k <= t]
            envs = [k for k in env_by_t if k <= t]
            if not hrs or not envs:
                continue
            a, rh = env_by_t[max(envs)]
            history.append({'time_min': t, 'hr_bpm': hr_by_t[max(hrs)],
                            'skin_c': skin_by_t[t], 'ambient_c': a, 'humidity_pct': rh})
        count = 0
        # Common cutoffs for all horizons avoids changing horizon weights near record end.
        for t in range(10, 41):
            if not history or not any(r['time_min'] == t for r in history) or any(t+h not in y for h in (5, 10, 20)):
                continue
            request = make_request(history, demo, t)
            records.append({'id': f'{person}-{t}', 'person': person, 'session': 'cycling',
                            'request': request, 'targets': [{'slice': h, 'value': y[t+h]} for h in (5, 10, 20)]})
            count += 1
        if not count:
            raise ValueError(f'No eligible forecasts for {person}')
        people.append(person)
    if len(people) != 24:
        raise ValueError('Expected 24 original people')
    sources = {'url': SOURCE, 'sha256': file_hash(path), 'license': 'CC BY 4.0', 'filename': path.name}
    return export_dataset(output_root, 'heat', records, people, sources,
                          {'sources': sources, 'flags': flags, 'lineage': lineage,
                           'time_rule': 'Current minute and interval right edge available; carry past HR/environment only',
                           'skin_channel': 'YSI TSK reference skin thermistor; not vendor CORE estimate',
                           'cutoffs_min': [10, 40], 'horizons_min': [5,10,20],
                           'excluded_inputs': ['rectal history','CORE','CORE TSK','VO2','RPE','wattage','subcutaneous fat']}, split_manifest)
